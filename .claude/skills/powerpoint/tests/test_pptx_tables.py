# Copyright (c) Microsoft Corporation.
# SPDX-License-Identifier: MIT
"""Tests for pptx_tables module."""

import pytest
from pptx.util import Inches, Pt
from pptx_tables import add_table_element, extract_table


def _make_table(slide, elem, colors=None, typography=None):
    """Helper to create a table on a slide."""
    return add_table_element(slide, elem, colors or {}, typography or {})


class TestAddTableElement:
    """Tests for add_table_element."""

    def test_basic_table(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 8.0,
            "height": 3.0,
            "columns": [{"width": 4.0}, {"width": 4.0}],
            "rows": [
                {"cells": [{"text": "A1"}, {"text": "B1"}]},
                {"cells": [{"text": "A2"}, {"text": "B2"}]},
            ],
        }
        shape = _make_table(blank_slide, elem)
        table = shape.table
        assert len(table.rows) == 2
        assert len(table.columns) == 2
        assert table.cell(0, 0).text == "A1"
        assert table.cell(1, 1).text == "B2"

    def test_table_properties(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 8.0,
            "height": 3.0,
            "first_row": True,
            "horz_banding": True,
            "columns": [{"width": 4.0}, {"width": 4.0}],
            "rows": [
                {"cells": [{"text": "H1"}, {"text": "H2"}]},
                {"cells": [{"text": "D1"}, {"text": "D2"}]},
            ],
        }
        shape = _make_table(blank_slide, elem)
        assert shape.table.first_row is True
        assert shape.table.horz_banding is True

    def test_cell_formatting(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 8.0,
            "height": 2.0,
            "columns": [{"width": 8.0}],
            "rows": [
                {"cells": [{"text": "Bold", "font_bold": True, "font_size": 14}]},
            ],
        }
        shape = _make_table(blank_slide, elem)
        cell = shape.table.cell(0, 0)
        run = cell.text_frame.paragraphs[0].runs[0]
        assert run.font.bold is True
        assert run.font.size == Pt(14)

    def test_cell_fill(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 8.0,
            "height": 2.0,
            "columns": [{"width": 8.0}],
            "rows": [
                {"cells": [{"text": "Filled", "fill": "#0078D4"}]},
            ],
        }
        _make_table(blank_slide, elem)
        # Verify no crashes — fill is applied

    def test_column_widths(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 8.0,
            "height": 2.0,
            "columns": [{"width": 2.0}, {"width": 6.0}],
            "rows": [
                {"cells": [{"text": "A"}, {"text": "B"}]},
            ],
        }
        shape = _make_table(blank_slide, elem)
        table = shape.table
        assert table.columns[0].width == Inches(2.0)
        assert table.columns[1].width == Inches(6.0)

    def test_shape_name(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 4.0,
            "height": 2.0,
            "name": "MyTable",
            "columns": [{"width": 4.0}],
            "rows": [{"cells": [{"text": "X"}]}],
        }
        shape = _make_table(blank_slide, elem)
        assert shape.name == "MyTable"

    def test_vertical_anchor(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 4.0,
            "height": 2.0,
            "columns": [{"width": 4.0}],
            "rows": [
                {"cells": [{"text": "Mid", "vertical_anchor": "middle"}]},
            ],
        }
        shape = _make_table(blank_slide, elem)
        from pptx.enum.text import MSO_VERTICAL_ANCHOR

        assert shape.table.cell(0, 0).vertical_anchor == MSO_VERTICAL_ANCHOR.MIDDLE


class TestExtractTable:
    """Tests for extract_table."""

    def test_basic_roundtrip(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 8.0,
            "height": 3.0,
            "columns": [{"width": 4.0}, {"width": 4.0}],
            "rows": [
                {"cells": [{"text": "A1"}, {"text": "B1"}]},
                {"cells": [{"text": "A2"}, {"text": "B2"}]},
            ],
        }
        shape = _make_table(blank_slide, elem)
        result = extract_table(shape)
        assert result["type"] == "table"
        assert len(result["rows"]) == 2
        assert result["rows"][0]["cells"][0]["text"] == "A1"
        assert result["rows"][1]["cells"][1]["text"] == "B2"

    def test_position_roundtrip(self, blank_slide):
        elem = {
            "left": 2.0,
            "top": 3.0,
            "width": 6.0,
            "height": 2.5,
            "columns": [{"width": 6.0}],
            "rows": [{"cells": [{"text": "X"}]}],
        }
        shape = _make_table(blank_slide, elem)
        result = extract_table(shape)
        assert result["left"] == pytest.approx(2.0, abs=0.01)
        assert result["top"] == pytest.approx(3.0, abs=0.01)
        assert result["width"] == pytest.approx(6.0, abs=0.01)

    def test_table_properties_roundtrip(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 8.0,
            "height": 3.0,
            "first_row": True,
            "horz_banding": True,
            "columns": [{"width": 4.0}, {"width": 4.0}],
            "rows": [
                {"cells": [{"text": "H1"}, {"text": "H2"}]},
                {"cells": [{"text": "D1"}, {"text": "D2"}]},
            ],
        }
        shape = _make_table(blank_slide, elem)
        result = extract_table(shape)
        assert result["first_row"] is True
        assert result["horz_banding"] is True

    def test_font_info_extracted(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 8.0,
            "height": 2.0,
            "columns": [{"width": 8.0}],
            "rows": [
                {"cells": [{"text": "Bold", "font_bold": True}]},
            ],
        }
        shape = _make_table(blank_slide, elem)
        result = extract_table(shape)
        cell = result["rows"][0]["cells"][0]
        assert cell.get("font_bold") is True

    def test_column_widths_roundtrip(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 8.0,
            "height": 2.0,
            "columns": [{"width": 2.0}, {"width": 6.0}],
            "rows": [{"cells": [{"text": "A"}, {"text": "B"}]}],
        }
        shape = _make_table(blank_slide, elem)
        result = extract_table(shape)
        assert len(result["columns"]) == 2
        assert result["columns"][0]["width"] == pytest.approx(2.0, abs=0.01)
        assert result["columns"][1]["width"] == pytest.approx(6.0, abs=0.01)

    def test_name_roundtrip(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 4.0,
            "height": 2.0,
            "name": "TestTable",
            "columns": [{"width": 4.0}],
            "rows": [{"cells": [{"text": "X"}]}],
        }
        shape = _make_table(blank_slide, elem)
        result = extract_table(shape)
        assert result["name"] == "TestTable"

    def test_table_banding_properties(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 4.0,
            "height": 2.0,
            "first_row": True,
            "last_row": True,
            "first_col": True,
            "last_col": True,
            "horz_banding": True,
            "vert_banding": True,
            "columns": [{"width": 4.0}],
            "rows": [{"cells": [{"text": "X"}]}],
        }
        shape = _make_table(blank_slide, elem)
        result = extract_table(shape)
        assert result.get("first_row") is True
        assert result.get("horz_banding") is True

    def test_cell_fill(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 4.0,
            "height": 2.0,
            "columns": [{"width": 4.0}],
            "rows": [{"cells": [{"text": "Filled", "fill": "#0078D4"}]}],
        }
        shape = _make_table(blank_slide, elem)
        result = extract_table(shape)
        cell = result["rows"][0]["cells"][0]
        assert "fill" in cell

    def test_cell_vertical_anchor(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 4.0,
            "height": 2.0,
            "columns": [{"width": 4.0}],
            "rows": [{"cells": [{"text": "Mid", "vertical_anchor": "middle"}]}],
        }
        shape = _make_table(blank_slide, elem)
        result = extract_table(shape)
        cell = result["rows"][0]["cells"][0]
        assert cell.get("vertical_anchor") == "middle"

    def test_merge_roundtrip(self, blank_slide):
        elem = {
            "left": 1.0,
            "top": 1.0,
            "width": 8.0,
            "height": 2.0,
            "columns": [{"width": 4.0}, {"width": 4.0}],
            "rows": [
                {
                    "cells": [
                        {"text": "Merged", "merge_right": 1},
                        {"text": ""},
                    ]
                }
            ],
        }
        shape = _make_table(blank_slide, elem)
        result = extract_table(shape)
        cell0 = result["rows"][0]["cells"][0]
        assert cell0.get("merge_right") == 1
