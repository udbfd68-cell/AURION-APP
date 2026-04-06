# Copyright (c) Microsoft Corporation.
# SPDX-License-Identifier: MIT
"""Tests for pptx_fills module."""

import pytest
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from pptx_fills import (
    DASH_STYLE_MAP,
    DASH_STYLE_REVERSE,
    apply_effect_list,
    apply_fill,
    apply_line,
    extract_effect_list,
    extract_fill,
    extract_line,
)


def _make_shape(slide):
    """Create a rectangle shape on a slide."""
    return slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2)
    )


# ----- apply_fill / extract_fill -----


class TestSolidFill:
    """Tests for solid fill apply and extract."""

    def test_string_color(self, blank_slide):
        shape = _make_shape(blank_slide)
        apply_fill(shape, "#0078D4", {})
        result = extract_fill(shape.fill)
        assert result == "#0078D4"

    def test_dict_solid(self, blank_slide):
        shape = _make_shape(blank_slide)
        apply_fill(shape, {"type": "solid", "color": "#FF0000"}, {})
        result = extract_fill(shape.fill)
        # Result should be either the hex string or a dict with color
        if isinstance(result, str):
            assert result == "#FF0000"
        else:
            assert result["color"] == "#FF0000"

    def test_solid_with_alpha(self, blank_slide):
        shape = _make_shape(blank_slide)
        apply_fill(shape, {"type": "solid", "color": "#0078D4", "alpha": 50.0}, {})
        result = extract_fill(shape.fill)
        assert isinstance(result, dict)
        assert result["type"] == "solid"
        assert result["alpha"] == pytest.approx(50.0, abs=0.1)

    def test_none_fill(self, blank_slide):
        shape = _make_shape(blank_slide)
        apply_fill(shape, None, {})
        result = extract_fill(shape.fill)
        assert result is None


class TestGradientFill:
    """Tests for gradient fill apply and extract."""

    def test_basic_gradient(self, blank_slide):
        shape = _make_shape(blank_slide)
        fill_spec = {
            "type": "gradient",
            "angle": 90,
            "stops": [
                {"position": 0.0, "color": "#FF0000"},
                {"position": 1.0, "color": "#0000FF"},
            ],
        }
        apply_fill(shape, fill_spec, {})
        result = extract_fill(shape.fill)
        assert result["type"] == "gradient"
        assert len(result["stops"]) == 2
        assert result["stops"][0]["position"] == pytest.approx(0.0, abs=0.01)
        assert result["stops"][1]["position"] == pytest.approx(1.0, abs=0.01)

    def test_gradient_with_alpha(self, blank_slide):
        shape = _make_shape(blank_slide)
        fill_spec = {
            "type": "gradient",
            "angle": 45,
            "stops": [
                {"position": 0.0, "color": "#FF0000", "alpha": 75.0},
                {"position": 1.0, "color": "#0000FF", "alpha": 25.0},
            ],
        }
        apply_fill(shape, fill_spec, {})
        result = extract_fill(shape.fill)
        assert result["type"] == "gradient"
        assert result["stops"][0]["alpha"] == pytest.approx(75.0, abs=0.1)
        assert result["stops"][1]["alpha"] == pytest.approx(25.0, abs=0.1)

    def test_gradient_three_stops(self, blank_slide):
        shape = _make_shape(blank_slide)
        fill_spec = {
            "type": "gradient",
            "angle": 0,
            "stops": [
                {"position": 0.0, "color": "#FF0000"},
                {"position": 0.5, "color": "#00FF00"},
                {"position": 1.0, "color": "#0000FF"},
            ],
        }
        apply_fill(shape, fill_spec, {})
        result = extract_fill(shape.fill)
        assert result["type"] == "gradient"
        assert len(result["stops"]) == 3


class TestPatternFill:
    """Tests for pattern fill apply and extract."""

    def test_basic_pattern(self, blank_slide):
        shape = _make_shape(blank_slide)
        fill_spec = {
            "type": "pattern",
            "pattern": "cross",
            "fore_color": "#000000",
            "back_color": "#FFFFFF",
        }
        apply_fill(shape, fill_spec, {})
        result = extract_fill(shape.fill)
        assert result["type"] == "pattern"
        assert "fore_color" in result
        assert "back_color" in result

    def test_pattern_with_alpha(self, blank_slide):
        shape = _make_shape(blank_slide)
        fill_spec = {
            "type": "pattern",
            "pattern": "cross",
            "fore_color": "#000000",
            "back_color": "#FFFFFF",
            "fore_alpha": 80.0,
            "back_alpha": 60.0,
        }
        apply_fill(shape, fill_spec, {})
        result = extract_fill(shape.fill)
        assert result["type"] == "pattern"
        assert result["fore_alpha"] == pytest.approx(80.0, abs=0.1)
        assert result["back_alpha"] == pytest.approx(60.0, abs=0.1)


class TestNonDictFill:
    """Tests for edge cases in apply_fill."""

    def test_non_dict_non_string(self, blank_slide):
        shape = _make_shape(blank_slide)
        apply_fill(shape, 42, {})
        # No crash — non-dict, non-str, non-None is a no-op


# ----- apply_line / extract_line -----


class TestLine:
    """Tests for apply_line and extract_line."""

    def test_line_color_and_width(self, blank_slide):
        shape = _make_shape(blank_slide)
        elem = {"line_color": "#FF0000", "line_width": 2}
        apply_line(shape, elem, {})
        result = extract_line(shape)
        assert result["line_color"] == "#FF0000"
        assert result["line_width"] == pytest.approx(2.0, abs=0.1)

    def test_line_dash_style(self, blank_slide):
        shape = _make_shape(blank_slide)
        elem = {"line_color": "#000000", "line_width": 1, "dash_style": "dash"}
        apply_line(shape, elem, {})
        result = extract_line(shape)
        assert result["dash_style"] == "dash"

    def test_no_line_color_sets_background(self, blank_slide):
        shape = _make_shape(blank_slide)
        elem = {}
        apply_line(shape, elem, {})
        # Line should be set to background (no color)
        result = extract_line(shape)
        assert "line_color" not in result or result.get("line_color") is None

    def test_extract_line_default(self, blank_slide):
        shape = _make_shape(blank_slide)
        result = extract_line(shape)
        assert isinstance(result, dict)


# ----- apply_effect_list / extract_effect_list -----


class TestEffectList:
    """Tests for apply_effect_list and extract_effect_list."""

    def test_outer_shadow_preset_roundtrip(self, blank_slide):
        shape = _make_shape(blank_slide)
        effect = {
            "type": "outer_shadow",
            "blurRad": "50800",
            "dist": "38100",
            "dir": "2700000",
            "algn": "tl",
            "rotWithShape": "0",
            "color": "black",
            "color_type": "preset",
        }
        apply_effect_list(shape, effect)
        result = extract_effect_list(shape)
        assert result["type"] == "outer_shadow"
        assert result["blurRad"] == "50800"
        assert result["dist"] == "38100"
        assert result["dir"] == "2700000"
        assert result["algn"] == "tl"
        assert result["color"] == "black"
        assert result["color_type"] == "preset"

    def test_outer_shadow_rgb_roundtrip(self, blank_slide):
        shape = _make_shape(blank_slide)
        effect = {
            "type": "outer_shadow",
            "blurRad": "40000",
            "color": "#FF0000",
            "color_type": "rgb",
        }
        apply_effect_list(shape, effect)
        result = extract_effect_list(shape)
        assert result["color"] == "#FF0000"
        assert result["color_type"] == "rgb"

    def test_outer_shadow_alpha(self, blank_slide):
        shape = _make_shape(blank_slide)
        effect = {
            "type": "outer_shadow",
            "blurRad": "40000",
            "color": "black",
            "color_type": "preset",
            "alpha": 50.0,
        }
        apply_effect_list(shape, effect)
        result = extract_effect_list(shape)
        assert result["alpha"] == pytest.approx(50.0, abs=0.1)

    def test_no_effect_list(self, blank_slide):
        shape = _make_shape(blank_slide)
        result = extract_effect_list(shape)
        assert result is None

    def test_non_shadow_type_noop(self, blank_slide):
        shape = _make_shape(blank_slide)
        apply_effect_list(shape, {"type": "glow"})
        assert extract_effect_list(shape) is None

    def test_none_effect_noop(self, blank_slide):
        shape = _make_shape(blank_slide)
        apply_effect_list(shape, None)
        assert extract_effect_list(shape) is None

    def test_replaces_existing(self, blank_slide):
        shape = _make_shape(blank_slide)
        effect1 = {
            "type": "outer_shadow",
            "blurRad": "40000",
            "color": "black",
            "color_type": "preset",
        }
        apply_effect_list(shape, effect1)
        effect2 = {
            "type": "outer_shadow",
            "blurRad": "80000",
            "color": "#FF0000",
            "color_type": "rgb",
        }
        apply_effect_list(shape, effect2)
        result = extract_effect_list(shape)
        assert result["blurRad"] == "80000"
        assert result["color"] == "#FF0000"


# ----- Constants -----


class TestFillConstants:
    """Tests for fill module constants."""

    def test_dash_style_map_keys(self):
        expected = {
            "solid",
            "dash",
            "dash_dot",
            "dash_dot_dot",
            "long_dash",
            "long_dash_dot",
            "round_dot",
            "square_dot",
        }
        assert set(DASH_STYLE_MAP.keys()) == expected

    def test_dash_style_reverse_consistency(self):
        for name, val in DASH_STYLE_MAP.items():
            assert DASH_STYLE_REVERSE[val] == name
